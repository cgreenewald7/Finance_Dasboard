"""Exercise the source and frozen app using disposable data, never a real budget."""
import csv
from pathlib import Path
import tempfile
from unittest.mock import patch


def run_checks():
    from budget_track import BudgetTracker
    from tkcalendar import Calendar
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory(prefix='budgettracker-check-') as folder:
        folder = Path(folder)
        tracker = BudgetTracker(data_dir=folder)
        try:
            def callback_error(exc_type, value, traceback):
                raise value.with_traceback(traceback)

            tracker.root.report_callback_exception = callback_error
            tracker.root.withdraw()
            tracker.current_month = '2026-01'
            tracker.all_income = [dict(source='Salary', amount=2000.0, date='2026-01-01', month='2026-01')]
            tracker.all_expenses = [dict(where='Café', amount=100.0, date='2026-01-02', month='2026-01',
                                         category='Groceries', payment_method='Credit Card', authorized_user='Alex')]
            tracker.all_savings = [dict(account='Retirement', amount=250.0, date='2026-01-03',
                                        month='2026-01', category='Retirement')]
            tracker.add_category('Custom')
            tracker.category_goals['Custom'] = 10.0
            tracker.category_goal_types['Custom'] = 'percent'
            tracker.deleted_categories.add('Activities')
            tracker.categories.remove('Activities')
            tracker.save_data()
            tracker.load_data()
            assert tracker.all_expenses[0]['authorized_user'] == 'Alex'
            assert tracker.all_expenses[0]['where'] == 'Café'
            assert tracker.all_income[0]['amount'] == 2000.0
            assert tracker.all_savings[0]['amount'] == 250.0
            assert tracker.category_goals['Custom'] == 10.0
            assert tracker.category_goal_types['Custom'] == 'percent'
            assert 'Activities' not in tracker.categories
            assert not tracker.get_month_data('2026-02')[0]
            assert tracker.get_savings_investing_total(tracker.all_expenses, tracker.all_savings) == 250.0
            Calendar(tracker.root, year=2026, month=1, day=1).destroy()
            tracker.toggle_theme()
            tracker.root.update()
            tracker.toggle_theme()
            tracker.root.update()
            csv_path = folder / 'export.csv'
            excel_path = folder / 'export.xlsx'
            with patch('budget_track.filedialog.asksaveasfilename', return_value=str(csv_path)), patch('budget_track.messagebox.showinfo'):
                tracker.export_selected_month()
            with csv_path.open(encoding='utf-8-sig', newline='') as handle:
                rows = list(csv.DictReader(handle))
            assert len(rows) == 3
            assert any(row['Authorized User'] == 'Alex' for row in rows)
            with patch('budget_track.filedialog.asksaveasfilename', return_value=str(excel_path)), patch('budget_track.messagebox.showinfo'):
                tracker.export_selected_month_excel()
            workbook = load_workbook(excel_path)
            assert workbook.active.max_row > 3
            workbook.close()
        finally:
            tracker.on_closing()
    print('Release checks passed: GUI, calendar, themes, persistence, goals, CSV and Excel.')


if __name__ == '__main__':
    run_checks()
